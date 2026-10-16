"""Word rendering (python-docx): the same report as the PDF, editable."""
import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from aura.core.domain.export import DealExportBundle, ExportTable
from aura.exports.common import RenderedExport, human_cell, human_datetime, humanize

INK = RGBColor(0x14, 0x17, 0x1A)
MUTED = RGBColor(0x67, 0x6E, 0x75)
DANGER = RGBColor(0xB4, 0x23, 0x18)
HEAD_FILL = "F4F3EF"
LINE = "E6E4DF"
NUMERIC = ("inr", "int", "float", "pct")
DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _field(run, code: str) -> None:
    """A Word field (PAGE, NUMPAGES) in a run."""
    begin, instr, end = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), \
        OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if edge in ("top", "bottom", "insideH") else "nil")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), LINE)
        borders.append(el)
    tbl_pr.append(borders)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _small(paragraph, text: str, size: float = 8, color=INK, bold: bool = False):
    run = paragraph.add_run(text)
    run.font.size, run.font.color.rgb, run.bold = Pt(size), color, bold
    return run


def _table(doc, table: ExportTable, width_cm: float) -> None:
    h = doc.add_paragraph()
    h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(8), Pt(3)
    _small(h, table.title, 9.5, bold=True)
    if table.note:
        _small(doc.add_paragraph(), table.note, 8, MUTED)
    if not table.rows:
        _small(doc.add_paragraph(), "None recorded.", 8, MUTED)
        return
    cols = [c for c in table.columns if c.human]
    t = doc.add_table(rows=1, cols=len(cols))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    _borders(t)
    total = sum(c.width for c in cols) or 1
    widths = [Cm(width_cm * c.width / total) for c in cols]
    head = t.rows[0]
    _repeat_header(head)
    for i, col in enumerate(cols):
        cell = head.cells[i]
        cell.width = widths[i]
        _shade(cell, HEAD_FILL)
        p = cell.paragraphs[0]
        if col.type in NUMERIC:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _small(p, col.label, 7.5, MUTED, bold=True)
    for row in table.rows:
        cells = t.add_row().cells
        for i, col in enumerate(cols):
            cells[i].width = widths[i]
            p = cells[i].paragraphs[0]
            if col.type in NUMERIC:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _small(p, human_cell(row.get(col.key), col.type), 7.5)
    if table.truncated:
        _small(doc.add_paragraph(), "Truncated: more rows exist than an export carries.", 7.5,
               MUTED)


def _facts(doc, facts: list[tuple[str, str]], width_cm: float) -> None:
    if not facts:
        return
    cols = 2
    t = doc.add_table(rows=0, cols=cols * 2)
    t.autofit = False
    for i in range(0, len(facts), cols):
        cells = t.add_row().cells
        for j, (label, value) in enumerate(facts[i:i + cols]):
            cells[j * 2].width, cells[j * 2 + 1].width = Cm(width_cm * 0.17), Cm(width_cm * 0.33)
            _small(cells[j * 2].paragraphs[0], label, 8, MUTED)
            _small(cells[j * 2 + 1].paragraphs[0], value, 8.5)


def render_docx(bundle: DealExportBundle) -> RenderedExport:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    width_cm = 21.0 - 3.6
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(9)
    d = bundle.deal

    head = sec.header.paragraphs[0]
    _small(head, f"{d['external_id']} · {d['name']}", 7.5, MUTED)
    _small(head, "    CONFIDENTIAL" + ("  ·  CONTAINS UPSI" if bundle.upsi else ""), 7.5,
           DANGER if bundle.upsi else MUTED, bold=True)
    foot = sec.footer.paragraphs[0]
    _small(foot, f"Generated {human_datetime(bundle.generated_at)} IST by "
                 f"{bundle.generated_by}  ·  Page ", 7.5, MUTED)
    _field(foot.add_run(), "PAGE")
    _small(foot, " of ", 7.5, MUTED)
    _field(foot.add_run(), "NUMPAGES")

    kicker = doc.add_paragraph()
    _small(kicker, f"{d['external_id']}  ·  {d['transaction_class']}  ·  "
                   f"{humanize(d['stage'])}", 9, MUTED)
    title = doc.add_paragraph()
    run = title.add_run(d["name"])
    run.font.size, run.font.name, run.font.color.rgb = Pt(24), "Georgia", INK
    sub = doc.add_paragraph()
    _small(sub, "Deal report", 10, MUTED)
    if bundle.upsi:
        warn = doc.add_paragraph()
        _small(warn, "Contains unpublished price-sensitive information (UPSI). Handle under "
                     "the insider-trading code; do not forward outside the deal wall.", 8.5,
               DANGER, bold=True)

    for section in bundle.sections:
        h = doc.add_paragraph()
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(14), Pt(4)
        r = h.add_run(section.title)
        r.font.size, r.font.name, r.font.color.rgb = Pt(15), "Georgia", INK
        _facts(doc, section.facts, width_cm)
        for table in section.tables:
            _table(doc, table, width_cm)
    doc.core_properties.title = f"{d['external_id']} {d['name']}"
    doc.core_properties.author = bundle.generated_by
    doc.core_properties.comments = "Generated by Aura. Confidential."
    out = io.BytesIO()
    doc.save(out)
    return RenderedExport(out.getvalue(), DOCX_MEDIA, "docx")
