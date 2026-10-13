"""Document parsers: the DocumentParser port for every format Aura reads.

    PDF   Docling (aura.ingest.docling_parser: layout, table cell grids, OCR) when installed
          and AURA_PDF_PARSER allows it; otherwise, or when Docling fails, times out or the
          PDF is over its page cap, PyMuPDF: per-page text with headings by size or bold, and
          pages with no text layer read by OCR (RapidOCR, local ONNX models; the render is
          deskewed and the boxes grouped into lines) up to ocr_max_pages. Either way, scanned
          pages can get a second reading by a local vision model (aura.ingest.vlm).
    DOCX  python-docx: paragraphs, headings (Heading/Title styles) and tables, in order
    XLSX  openpyxl: each sheet's used range as a table (header row kept)
    CSV   csv: one table
    PPTX  python-pptx: per slide, title, text, tables and speaker notes
    EML   email (stdlib): headers, plain or HTML body, attachments listed and parsed
    MSG   extract-msg: as EML
    TXT/MD  plain text; Markdown '#' lines are headings

The format is chosen by extension and checked against the file's magic bytes. A file named
.pdf (or .docx, ...) whose content is plain text is read as text, with a warning. Anything
unreadable raises a typed IngestError (UnsupportedFormat, CorruptDocument, DocumentTooLarge),
never a bare exception.
"""
import csv
import email
import html
import io
import logging
import re
import threading
import zipfile
from dataclasses import dataclass
from datetime import datetime
from email import policy
from html.parser import HTMLParser
from pathlib import PurePath

from aura.core.domain.ingest import (Block, CorruptDocument, DocumentTooLarge, IngestError,
                                     ParsedDocument, UnsupportedFormat)

logger = logging.getLogger(__name__)

FORMATS = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".xlsm": "xlsx", ".csv": "csv",
           ".pptx": "pptx", ".eml": "eml", ".msg": "msg", ".txt": "txt", ".md": "md",
           ".markdown": "md"}
MIME = {"application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
        "text/csv": "csv", "message/rfc822": "eml", "application/vnd.ms-outlook": "msg",
        "text/plain": "txt", "text/markdown": "md"}
_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
MAX_SHEET_ROWS = 20_000
MAX_ATTACHMENT_DEPTH = 2


@dataclass(frozen=True)
class ParserLimits:
    """Limits and PDF routing. The defaults are the plain PyMuPDF path with no vision model
    (what unit tests want); limits_from_settings() gives the configured ones."""
    max_upload_mb: int = 25
    max_pages: int = 300
    ocr_enabled: bool = True
    ocr_max_pages: int = 50
    ocr_min_chars: int = 25
    pdf_parser: str = "pymupdf"  # pymupdf | docling (settings' 'auto' is resolved here)
    docling_artifacts: str | None = None
    docling_max_pages: int = 120
    docling_timeout_s: float = 600.0
    scan_vlm: str = "off"  # off | fallback | always
    vlm_max_pages: int = 5

    @property
    def max_bytes(self) -> int:
        return self.max_upload_mb * 1024 * 1024


def limits_from_settings() -> ParserLimits:
    from aura.config import get_settings
    s = get_settings()
    return ParserLimits(max_upload_mb=s.max_upload_mb, max_pages=s.max_pages,
                        ocr_enabled=s.ocr_enabled, ocr_max_pages=s.ocr_max_pages,
                        ocr_min_chars=s.ocr_min_chars,
                        pdf_parser=resolve_pdf_parser(s.pdf_parser, s.docling_artifacts),
                        docling_artifacts=s.docling_artifacts,
                        docling_max_pages=s.docling_max_pages,
                        docling_timeout_s=s.docling_timeout_s, scan_vlm=s.scan_vlm,
                        vlm_max_pages=s.vlm_max_pages)


def resolve_pdf_parser(choice: str, artifacts: str | None) -> str:
    """'auto' -> 'docling' when the optional dependency and its models are installed, else
    'pymupdf'. An explicit 'docling' stays (a missing install then falls back per document,
    with a warning saying why)."""
    if choice != "auto":
        return choice
    from aura.ingest.docling_parser import docling_available
    return "docling" if docling_available(artifacts) else "pymupdf"


def extension(filename: str) -> str:
    return PurePath(filename or "").suffix.lower()


def supported_extensions() -> list[str]:
    return sorted(FORMATS)


def _looks_like_text(content: bytes) -> bool:
    head = content[:4096]
    if not head or b"\x00" in head:
        return False
    try:
        head.decode("utf-8")
        return True
    except UnicodeDecodeError as exc:  # a multi-byte character cut at the boundary is fine
        if exc.start >= len(head) - 4:
            return True
    printable = sum(1 for b in head if b in (9, 10, 13) or 32 <= b < 127 or b >= 160)
    return printable / len(head) > 0.95


def sniff(content: bytes) -> str | None:
    """The format the bytes look like, or None."""
    if content.startswith(b"%PDF"):
        return "pdf"
    if content.startswith(b"PK\x03\x04"):
        try:
            names = zipfile.ZipFile(io.BytesIO(content)).namelist()
        except zipfile.BadZipFile:
            return None
        for prefix, kind in (("word/", "docx"), ("xl/", "xlsx"), ("ppt/", "pptx")):
            if any(n.startswith(prefix) for n in names):
                return kind
        return None
    if content.startswith(_OLE):
        return "msg"
    if _looks_like_text(content):
        head = content[:2048].decode("utf-8", errors="replace")
        if re.search(r"^(From|Received|Message-ID|MIME-Version|Subject|Return-Path):",
                     head, re.MULTILINE | re.IGNORECASE) and "\n\n" in \
                content[:65536].decode("utf-8", errors="replace").replace("\r\n", "\n"):
            return "eml"
        return "txt"
    return None


def detect(content: bytes, filename: str, mime: str | None = None) -> tuple[str, list[str]]:
    """(format, warnings). Extension first, checked against the magic bytes."""
    ext = extension(filename)
    claimed = FORMATS.get(ext) or MIME.get((mime or "").split(";")[0].strip().lower())
    actual = sniff(content)
    if claimed is None:
        if ext and ext not in ("",):
            raise UnsupportedFormat(
                f"{filename}: {ext} files are not supported (supported: "
                f"{', '.join(supported_extensions())})")
        if actual is None:
            raise UnsupportedFormat(f"{filename}: unrecognised file type")
        return actual, []
    if claimed in ("txt", "md", "csv", "eml"):
        if actual not in ("txt", "eml"):
            raise CorruptDocument(f"{filename}: not a text file")
        return claimed, []
    if actual == claimed:
        return claimed, []
    if actual in ("txt", "eml"):
        return actual, [f"the content is plain text, not {claimed.upper()}; read as text"]
    raise CorruptDocument(f"{filename}: the content is not a valid {claimed.upper()} file")


def check_upload(filename: str, content: bytes, limits: ParserLimits | None = None) -> str:
    """Refuse an upload Aura could never index: unsupported type (415) or too large (413).
    Returns the format. Content problems beyond the type are left to the indexer."""
    limits = limits or limits_from_settings()
    if len(content) > limits.max_bytes:
        raise DocumentTooLarge(f"{filename}: {len(content) / 1048576:.1f} MB is over the "
                               f"{limits.max_upload_mb} MB upload limit")
    ext = extension(filename)
    if ext and ext not in FORMATS:
        raise UnsupportedFormat(f"{ext} files are not supported (supported: "
                                f"{', '.join(supported_extensions())})")
    if not ext and sniff(content) is None:
        raise UnsupportedFormat(f"{filename}: unrecognised file type")
    return FORMATS.get(ext) or sniff(content) or "txt"


def decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.6g}" \
            if abs(value) < 1e15 else str(value)
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() \
            else value.isoformat(sep=" ")
    return " ".join(str(value).split())


def _trim_table(rows: list[list[str]]) -> list[list[str]]:
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return []
    width = max(max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows)
    lead = min(next((i for i, c in enumerate(r) if c), width) for r in rows)
    return [(r + [""] * width)[lead:width] for r in rows]


def _table_block(rows: list[list[str]], page: int | None = None) -> Block:
    rows = _trim_table(rows)
    return Block(kind="table", text="\n".join(" | ".join(r) for r in rows), page=page,
                 rows=rows)


def _paragraph_blocks(text: str, page: int | None = None, markdown: bool = False
                      ) -> list[Block]:
    blocks: list[Block] = []
    for para in re.split(r"\n\s*\n", _clean(text)):
        para = para.strip()
        if not para:
            continue
        if markdown:
            lines, body = para.split("\n"), []
            for line in lines:
                m = re.match(r"^(#{1,6})\s+(.*)$", line)
                if m:
                    if body:
                        blocks.append(Block("text", "\n".join(body), page))
                        body = []
                    blocks.append(Block("heading", m.group(2).strip(), page,
                                        level=len(m.group(1))))
                else:
                    body.append(line)
            if body:
                blocks.append(Block("text", "\n".join(body), page))
        else:
            blocks.append(Block("text", para, page))
    return blocks


# ------------------------------------------------------------------ OCR

class _Ocr:
    """RapidOCR (ONNX models shipped in the wheel), loaded once, one page at a time."""
    _lock = threading.Lock()
    _engine = None
    _failed: str | None = None

    @classmethod
    def boxes(cls, png: bytes) -> list:
        """RapidOCR's raw result: [[box (4 points), text, score], ...]."""
        with cls._lock:
            if cls._engine is None and cls._failed is None:
                try:
                    from rapidocr_onnxruntime import RapidOCR
                    cls._engine = RapidOCR()
                except Exception as exc:  # pragma: no cover - depends on the machine
                    cls._failed = f"{type(exc).__name__}: {exc}"
            if cls._engine is None:
                raise RuntimeError(f"OCR unavailable ({cls._failed})")
            result, _ = cls._engine(png)
        return result or []

    @classmethod
    def read(cls, png: bytes) -> str:
        return ocr_lines_text(cls.boxes(png))


def _box(entry) -> tuple[float, float, float, float, str] | None:
    try:
        points, text, score = entry[0], entry[1], float(entry[2])
    except (TypeError, ValueError, IndexError):
        return None
    text = (text or "").strip()
    if not text or score < 0.5:
        return None
    xs, ys = [float(p[0]) for p in points], [float(p[1]) for p in points]
    return min(xs), min(ys), max(xs), max(ys), text


def ocr_lines(result: list) -> list[list[tuple]]:
    """Group OCR boxes into text lines by vertical overlap (not fixed-height buckets, which
    scatter the rows of a slightly rotated table), each line left to right. A box joins the
    line it overlaps by at least half its height (the line's mean top and bottom)."""
    boxes = sorted((b for b in (_box(e) for e in result) if b),
                   key=lambda b: ((b[1] + b[3]) / 2, b[0]))
    lines: list[dict] = []
    for b in boxes:
        height = max(1.0, b[3] - b[1])
        best, best_overlap = None, 0.0
        for line in lines[-4:]:
            top, bottom = line["top"] / line["n"], line["bottom"] / line["n"]
            overlap = min(bottom, b[3]) - max(top, b[1])
            if overlap >= 0.5 * min(height, bottom - top) and overlap > best_overlap:
                best, best_overlap = line, overlap
        if best is None:
            best = {"top": 0.0, "bottom": 0.0, "n": 0, "boxes": []}
            lines.append(best)
        best["boxes"].append(b)
        best["top"] += b[1]
        best["bottom"] += b[3]
        best["n"] += 1
    lines.sort(key=lambda ln: (ln["top"] + ln["bottom"]) / (2 * ln["n"]))
    return [sorted(ln["boxes"], key=lambda b: b[0]) for ln in lines]


def _join_line(boxes: list[tuple]) -> str:
    """Boxes of one line joined left to right: a space between boxes, none where two boxes
    touch (one word cut in two)."""
    out = ""
    prev = None
    for b in boxes:
        if prev is not None:
            char_w = max(1.0, (prev[2] - prev[0]) / max(1, len(prev[4])))
            out += "" if b[0] - prev[2] < 0.25 * char_w else " "
        out += b[4]
        prev = b
    return out


def ocr_lines_text(result: list) -> str:
    """OCR boxes as text: one line per text line, a blank line where the vertical gap is
    larger than one and a half line heights (a paragraph break)."""
    lines = ocr_lines(result)
    out: list[str] = []
    prev_bottom = None
    heights = sorted(max(1.0, max(b[3] for b in ln) - min(b[1] for b in ln)) for ln in lines)
    typical = heights[len(heights) // 2] if heights else 0.0
    for ln in lines:
        top, bottom = min(b[1] for b in ln), max(b[3] for b in ln)
        if prev_bottom is not None and top - prev_bottom > 1.5 * typical:
            out.append("")
        out.append(_join_line(ln))
        prev_bottom = bottom
    return "\n".join(out)


def estimate_skew(gray, max_angle: float = 5.0) -> float:
    """Skew of a page image in degrees (positive: text lines fall to the right), by the
    projection profile: the shear that makes the row sums of ink most peaked. numpy only."""
    import numpy as np
    ink = np.asarray(gray) < 128
    ys, xs = np.nonzero(ink)
    if len(ys) < 200:
        return 0.0
    if len(ys) > 400_000:  # plenty for a profile; keep it fast
        pick = np.random.default_rng(0).choice(len(ys), 400_000, replace=False)
        ys, xs = ys[pick], xs[pick]
    xs = xs - xs.mean()

    def score(angle: float) -> float:
        yy = np.round(ys - xs * np.tan(np.radians(angle))).astype(np.int64)
        hist = np.bincount(yy - yy.min())
        return float(np.square(hist.astype(np.float64)).sum())

    coarse = np.arange(-max_angle, max_angle + 1e-9, 0.25)
    best = max(coarse, key=score)
    fine = np.arange(best - 0.25, best + 0.25 + 1e-9, 0.05)
    return round(float(max(fine, key=score)), 2)


def render_for_ocr(page, dpi: int = 200, min_skew: float = 0.3) -> tuple[bytes, float]:
    """The page as PNG for OCR, straightened when the scan is skewed (more than min_skew
    degrees). Returns (png, angle corrected)."""
    import fitz
    import numpy as np
    angle = 0.0
    try:
        probe = page.get_pixmap(dpi=100, colorspace=fitz.csGRAY)
        gray = np.frombuffer(probe.samples, dtype=np.uint8).reshape(probe.height,
                                                                     probe.stride)
        angle = estimate_skew(gray[:, :probe.width])
    except Exception:  # pragma: no cover - a page MuPDF cannot render is read unrotated
        angle = 0.0
    zoom = dpi / 72.0
    if abs(angle) >= min_skew:
        matrix = fitz.Matrix(zoom, zoom).prerotate(-angle)
        return page.get_pixmap(matrix=matrix).tobytes("png"), angle
    return page.get_pixmap(dpi=dpi).tobytes("png"), 0.0


def ocr_available() -> bool:
    try:
        import rapidocr_onnxruntime  # noqa: F401
        return True
    except Exception:
        return False


# ------------------------------------------------------------------ formats

_BOLD = 16  # PyMuPDF span flag


def _open_pdf(content: bytes, limits: ParserLimits):
    import fitz  # PyMuPDF
    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise CorruptDocument(f"the PDF cannot be opened ({exc})") from exc
    if doc.needs_pass:
        doc.close()
        raise CorruptDocument("the PDF is password-protected")
    if doc.page_count > limits.max_pages:
        count = doc.page_count
        doc.close()
        raise DocumentTooLarge(f"{count} pages is over the {limits.max_pages}-page limit")
    return doc


def image_only_pages(doc, min_chars: int) -> list[int]:
    """Pages (1-based) with fewer than min_chars characters in their text layer."""
    out = []
    for page in doc:
        try:
            chars = len(page.get_text("text").strip())
        except Exception:
            chars = 0
        if chars < min_chars:
            out.append(page.number + 1)
    return out


def _heading_level(size: float, body: float) -> int:
    ratio = size / body if body else 1.0
    return 1 if ratio >= 1.5 else 2 if ratio >= 1.25 else 3


def parse_pdf(content: bytes, limits: ParserLimits) -> ParsedDocument:
    with _open_pdf(content, limits) as doc:
        return _parse_pdf_pymupdf(doc, limits)


def _parse_pdf_pymupdf(doc, limits: ParserLimits) -> ParsedDocument:
    warnings: list[str] = []
    pages = []
    sizes: dict[float, int] = {}
    for page in doc:
        try:
            data = page.get_text("dict", flags=0)
        except Exception as exc:
            warnings.append(f"page {page.number + 1}: text could not be read ({exc})")
            data = {"blocks": []}
        found = []  # per text block: [(line text, size, bold, line bbox), ...]
        for block in data.get("blocks", []):
            if block.get("type", 0) != 0:
                continue
            lines = []
            for line in block.get("lines", []):
                spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
                if not spans:
                    continue
                size, bold = 0.0, True
                for s in spans:
                    rounded = round(s.get("size", 0), 1)
                    sizes[rounded] = sizes.get(rounded, 0) + len(s["text"])
                    size = max(size, rounded)
                    bold = bold and bool(s.get("flags", 0) & _BOLD or
                                         "bold" in str(s.get("font", "")).lower())
                lines.append(("".join(s["text"] for s in line["spans"]).strip(), size, bold,
                              tuple(line.get("bbox") or (0, 0, 0, 0))))
            if lines:
                found.append(lines)
        pages.append(found)
    body_size = max(sizes, key=lambda k: sizes[k]) if sizes else 0
    blocks: list[Block] = []
    ocr_pages = 0
    for number, found in enumerate(pages, start=1):
        chars = sum(len(ln[0]) for lines in found for ln in lines)
        if chars < limits.ocr_min_chars:
            if not limits.ocr_enabled:
                warnings.append(f"page {number}: no text layer (OCR is off)")
            elif ocr_pages >= limits.ocr_max_pages:
                warnings.append(f"page {number}: no text layer; OCR limit of "
                                f"{limits.ocr_max_pages} pages reached")
            else:
                ocr_pages += 1
                try:
                    png, angle = render_for_ocr(doc[number - 1])
                    text = _Ocr.read(png)
                    if angle:
                        warnings.append(f"page {number}: scan straightened by "
                                        f"{angle:+.1f} degrees before OCR")
                except Exception as exc:
                    warnings.append(f"page {number}: OCR failed ({exc})")
                    text = ""
                if text.strip():
                    warnings.append(f"page {number}: text recognised by OCR")
                    for b in _paragraph_blocks(text, number):
                        b.meta["source"] = "ocr"
                        blocks.append(b)
                    continue
                if not found:
                    warnings.append(f"page {number}: no text found")
        blocks += _page_blocks(found, body_size, number)
    meta = {k: v for k, v in (doc.metadata or {}).items() if v}
    meta["parser"] = "pymupdf"
    return ParsedDocument(blocks=blocks, metadata=meta, page_count=doc.page_count,
                          ocr_used=ocr_pages > 0 and any("OCR" in w and "recognised" in w
                                                         for w in warnings),
                          warnings=warnings)


def _is_heading(line: tuple, body_size: float, others: list[tuple]) -> bool:
    """A line is a heading when it has words, is short, and is either set larger than the
    body (1.12x: a 10.5 pt heading over 9 pt text) or bold at body size while standing alone
    on its line (not a cell of a table row) and not reading as a sentence."""
    text, size, bold, bbox = line
    text = text.strip()
    if not text or len(text) > 120 or not re.search(r"[A-Za-z]{2}", text) or not body_size:
        return False
    if size >= body_size * 1.12:
        return True
    if not bold or len(text) > 100 or text.endswith((".", ",", ";")) or \
            size < body_size * 0.95:
        return False
    return not any(o is not bbox and o[1] < bbox[3] - 1 and o[3] > bbox[1] + 1 and
                   (o[2] < bbox[0] or o[0] > bbox[2]) for o in others)


def _page_blocks(found: list[list[tuple]], body_size: float, number: int) -> list[Block]:
    """A page's text blocks, with heading lines split out of them (a heading and its body
    often share one PyMuPDF block)."""
    others = [ln[3] for lines in found for ln in lines]
    out: list[Block] = []
    for lines in found:
        body: list[str] = []
        for i, line in enumerate(lines):
            # Only a block's leading lines can be headings: a bold word inside a paragraph
            # is emphasis.
            if not body and _is_heading(line, body_size, others):
                out.append(Block("heading", _clean(line[0]), number,
                                 level=_heading_level(line[1], body_size)))
            else:
                body.append(line[0])
        if body:
            out.append(Block("text", _clean("\n".join(body)), number))
    return out


def parse_pdf_routed(content: bytes, filename: str, limits: ParserLimits,
                     vlm=None) -> ParsedDocument:
    """A PDF through the configured parser and its fallback chain, then the vision-model pass
    over scanned pages:

        docling (if chosen) --missing / error / timeout / over docling_max_pages--> pymupdf
        image-only pages with thin OCR text (or a KYC / board-resolution scan) --> vision model

    Every fallback leaves a warning on the document saying why."""
    with _open_pdf(content, limits) as doc:
        warnings: list[str] = []
        parsed = None
        if limits.pdf_parser == "docling":
            from aura.ingest import docling_parser as dl
            reason = None
            if not dl.docling_available(limits.docling_artifacts):
                reason = dl.unavailable_reason(limits.docling_artifacts)
            elif doc.page_count > limits.docling_max_pages:
                reason = (f"{doc.page_count} pages is over the Docling cap of "
                          f"{limits.docling_max_pages}")
            else:
                try:
                    parsed = dl.parse_with_docling(content, filename, limits.docling_artifacts,
                                                   limits.docling_timeout_s, doc.page_count)
                    if not parsed.blocks and doc.page_count:
                        raise RuntimeError("Docling found no text")
                except Exception as exc:
                    logger.warning("Docling failed on %s; using PyMuPDF", filename,
                                   exc_info=True)
                    reason = f"Docling failed ({type(exc).__name__}: {str(exc)[:200]})"
                    parsed = None
            if parsed is None:
                warnings.append(f"{reason}; read with PyMuPDF instead")
        image_pages = image_only_pages(doc, limits.ocr_min_chars)
        if parsed is None:
            parsed = _parse_pdf_pymupdf(doc, limits)
        else:
            parsed.metadata.update({k: v for k, v in (doc.metadata or {}).items() if v})
            parsed.metadata["parser"] = "docling"
            ocr_pages = [p for p in image_pages
                         if any(b.page == p and (b.text or b.rows) for b in parsed.blocks)]
            parsed.ocr_used = bool(ocr_pages)
            if ocr_pages:
                warnings.append("page(s) " + ", ".join(map(str, ocr_pages)) +
                                ": text recognised by OCR (Docling)")
        parsed.warnings = warnings + parsed.warnings
        if limits.scan_vlm != "off" and image_pages:
            from aura.ingest import vlm as vlm_mod
            reader = vlm if vlm is not None else vlm_mod.vlm_from_settings()

            def page_png(number: int) -> bytes:
                return doc[number - 1].get_pixmap(dpi=vlm_mod.RENDER_DPI).tobytes("png")

            vlm_mod.apply_vlm(parsed, page_png, image_pages, limits.scan_vlm, filename,
                              limits.vlm_max_pages, reader)
        return parsed


def parse_docx(content: bytes, limits: ParserLimits) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise CorruptDocument(f"the DOCX cannot be opened ({exc})") from exc
    blocks: list[Block] = []
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, document)
            text = _clean(p.text)
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"(Heading|Title)\s*(\d*)", style)
            if m:
                blocks.append(Block("heading", text, None,
                                    level=int(m.group(2)) if m.group(2) else 1))
            else:
                blocks.append(Block("text", text))
        elif tag == "tbl":
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells, seen = [], set()
                for cell in row.cells:  # merged cells repeat: keep each once
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    cells.append(_cell(cell.text))
                rows.append(cells)
            block = _table_block(rows)
            if block.rows:
                blocks.append(block)
    props = document.core_properties
    meta = {k: getattr(props, k) for k in ("title", "author", "subject")
            if getattr(props, k, None)}
    return ParsedDocument(blocks=blocks, metadata=meta)


def parse_xlsx(content: bytes, limits: ParserLimits) -> ParsedDocument:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise CorruptDocument(f"the workbook cannot be opened ({exc})") from exc
    blocks: list[Block] = []
    warnings: list[str] = []
    try:
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_SHEET_ROWS:
                    warnings.append(f"sheet '{ws.title}': only the first {MAX_SHEET_ROWS} "
                                    "rows were indexed")
                    break
                rows.append([_cell(v) for v in row])
            table = _table_block(rows)
            if not table.rows:
                warnings.append(f"sheet '{ws.title}' is empty")
                continue
            blocks.append(Block("heading", f"Sheet: {ws.title}"))
            blocks.append(table)
    finally:
        wb.close()
    return ParsedDocument(blocks=blocks, metadata={"sheets": len(blocks) // 2},
                          warnings=warnings)


def parse_csv(content: bytes, limits: ParserLimits) -> ParsedDocument:
    text = decode_text(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[_cell(c) for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    warnings = []
    if len(rows) > MAX_SHEET_ROWS:
        rows = rows[:MAX_SHEET_ROWS]
        warnings.append(f"only the first {MAX_SHEET_ROWS} rows were indexed")
    table = _table_block(rows)
    return ParsedDocument(blocks=[table] if table.rows else [], warnings=warnings)


def parse_pptx(content: bytes, limits: ParserLimits) -> ParsedDocument:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    try:
        deck = Presentation(io.BytesIO(content))
    except Exception as exc:
        raise CorruptDocument(f"the presentation cannot be opened ({exc})") from exc
    slides = list(deck.slides)
    if len(slides) > limits.max_pages:
        raise DocumentTooLarge(f"{len(slides)} slides is over the {limits.max_pages} limit")
    blocks: list[Block] = []

    def shapes_of(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shapes_of(shape.shapes)
            else:
                yield shape

    for number, slide in enumerate(slides, start=1):
        title_shape = slide.shapes.title
        title = _clean(title_shape.text_frame.text) if title_shape is not None and \
            title_shape.has_text_frame else ""
        blocks.append(Block("heading", f"Slide {number}: {title}" if title
                            else f"Slide {number}", number))
        for shape in shapes_of(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[_cell(c.text) for c in r.cells] for r in shape.table.rows]
                block = _table_block(rows, number)
                if block.rows:
                    blocks.append(block)
            elif shape.has_text_frame:
                text = _clean(shape.text_frame.text)
                if text:
                    blocks.append(Block("text", text, number))
        if slide.has_notes_slide:
            notes = _clean(slide.notes_slide.notes_text_frame.text
                           if slide.notes_slide.notes_text_frame else "")
            if notes:
                blocks.append(Block("text", f"Speaker notes: {notes}", number))
    return ParsedDocument(blocks=blocks, page_count=len(slides),
                          metadata={"slides": len(slides)})


class _HtmlText(HTMLParser):
    _BREAK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table",
              "ul", "ol", "section", "article", "blockquote"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "head"):
            self._skip += 1
        elif tag in self._BREAK:
            self.parts.append("\n")
        elif tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "head"):
            self._skip = max(0, self._skip - 1)
        elif tag in self._BREAK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def html_to_text(markup: str) -> str:
    parser = _HtmlText()
    try:
        parser.feed(markup)
        parser.close()
    except Exception:
        return html.unescape(re.sub(r"<[^>]+>", " ", markup))
    return _clean("".join(parser.parts))


def _email_blocks(headers: list[tuple[str, str]], subject: str, body: str,
                  attachments: list[tuple[str, bytes]], limits: ParserLimits, depth: int,
                  vlm=None) -> ParsedDocument:
    blocks = [Block("heading", subject or "(no subject)")]
    head = "\n".join(f"{k}: {v}" for k, v in headers if v)
    if head:
        blocks.append(Block("text", head))
    blocks += _paragraph_blocks(body)
    warnings: list[str] = []
    ocr_used = False
    model_runs: list = []
    if attachments:
        listing = ", ".join(f"{name} ({max(1, len(data) // 1024)} KB)"
                            for name, data in attachments)
        blocks.append(Block("text", f"Attachments: {listing}"))
    for name, data in attachments:
        if depth >= MAX_ATTACHMENT_DEPTH:
            warnings.append(f"attachment {name}: nested too deep, not read")
            continue
        try:
            inner = parse(data, name, limits=limits, _depth=depth + 1, vlm=vlm)
        except UnsupportedFormat:
            warnings.append(f"attachment {name}: format not supported, listed only")
            continue
        except IngestError as exc:
            warnings.append(f"attachment {name}: {exc}")
            continue
        blocks.append(Block("heading", f"Attachment: {name}"))
        blocks += [Block(b.kind, b.text, None, b.rows, b.level, dict(b.meta))
                   for b in inner.blocks]
        warnings += [f"attachment {name}: {w}" for w in inner.warnings]
        ocr_used = ocr_used or inner.ocr_used
        model_runs += inner.model_runs
    return ParsedDocument(blocks=blocks, warnings=warnings, ocr_used=ocr_used,
                          metadata={"subject": subject, "attachments": len(attachments)},
                          model_runs=model_runs)


def parse_eml(content: bytes, limits: ParserLimits, depth: int = 0, vlm=None
              ) -> ParsedDocument:
    try:
        msg = email.message_from_bytes(content, policy=policy.default)
    except Exception as exc:
        raise CorruptDocument(f"the email cannot be read ({exc})") from exc
    headers = [(k, str(msg.get(k, "") or "")) for k in ("From", "To", "Cc", "Date",
                                                          "Subject")]
    body = ""
    try:
        part = msg.get_body(preferencelist=("plain", "html"))
        if part is not None:
            text = part.get_content()
            body = html_to_text(text) if part.get_content_type() == "text/html" else text
    except Exception:
        body = decode_text(msg.get_payload(decode=True) or b"") if not msg.is_multipart() \
            else ""
    attachments = []
    for att in msg.iter_attachments():
        name = att.get_filename() or "attachment"
        try:
            data = att.get_payload(decode=True) or b""
        except Exception:
            data = b""
        attachments.append((name, data))
    return _email_blocks(headers, str(msg.get("Subject", "") or ""), body, attachments,
                         limits, depth, vlm)


def parse_msg(content: bytes, limits: ParserLimits, depth: int = 0, vlm=None
              ) -> ParsedDocument:
    import extract_msg
    try:
        msg = extract_msg.Message(content)
    except Exception as exc:
        raise CorruptDocument(f"the Outlook message cannot be read ({exc})") from exc
    try:
        body = msg.body or ""
        if not body.strip() and getattr(msg, "htmlBody", None):
            raw = msg.htmlBody
            body = html_to_text(raw.decode("utf-8", errors="replace")
                                if isinstance(raw, bytes) else str(raw))
        headers = [("From", msg.sender or ""), ("To", msg.to or ""), ("Cc", msg.cc or ""),
                   ("Date", str(msg.date or "")), ("Subject", msg.subject or "")]
        attachments = []
        for att in msg.attachments:
            data = getattr(att, "data", None)
            if isinstance(data, bytes):
                name = getattr(att, "longFilename", None) or getattr(att, "shortFilename",
                                                                     None) or "attachment"
                attachments.append((name, data))
        return _email_blocks(headers, msg.subject or "", body, attachments, limits, depth,
                             vlm)
    finally:
        try:
            msg.close()
        except Exception:
            pass


def parse_text(content: bytes, limits: ParserLimits, markdown: bool = False
               ) -> ParsedDocument:
    return ParsedDocument(blocks=_paragraph_blocks(decode_text(content), None, markdown))


def parse(content: bytes, filename: str, mime: str | None = None,
          limits: ParserLimits | None = None, _depth: int = 0, vlm=None) -> ParsedDocument:
    limits = limits or limits_from_settings()
    if len(content) > limits.max_bytes:
        raise DocumentTooLarge(f"{len(content) / 1048576:.1f} MB is over the "
                               f"{limits.max_upload_mb} MB limit")
    if not content:
        raise CorruptDocument(f"{filename}: the file is empty")
    kind, warnings = detect(content, filename, mime)
    try:
        if kind == "pdf":
            parsed = parse_pdf_routed(content, filename, limits, vlm)
        elif kind == "docx":
            parsed = parse_docx(content, limits)
        elif kind == "xlsx":
            parsed = parse_xlsx(content, limits)
        elif kind == "csv":
            parsed = parse_csv(content, limits)
        elif kind == "pptx":
            parsed = parse_pptx(content, limits)
        elif kind == "eml":
            parsed = parse_eml(content, limits, _depth, vlm)
        elif kind == "msg":
            parsed = parse_msg(content, limits, _depth, vlm)
        else:
            parsed = parse_text(content, limits, markdown=kind == "md")
    except IngestError:
        raise
    except Exception as exc:  # a library failing on a damaged file is a corrupt document
        logger.warning("parser %s failed on %s", kind, filename, exc_info=True)
        raise CorruptDocument(f"{filename}: could not be read as {kind.upper()} "
                              f"({type(exc).__name__}: {exc})") from exc
    parsed.warnings = warnings + parsed.warnings
    parsed.metadata.setdefault("format", kind)
    return parsed


class FormatRegistryParser:
    """The DocumentParser port: format by extension and magic bytes, limits and PDF routing
    from settings (or the limits given). `vlm` replaces the vision model built from settings
    (tests pass a fake)."""

    def __init__(self, limits: ParserLimits | None = None, vlm=None):
        self.limits, self.vlm = limits, vlm

    def parse(self, content: bytes, filename: str, mime: str | None = None) -> ParsedDocument:
        return parse(content, filename, mime, self.limits or limits_from_settings(),
                     vlm=self.vlm)
