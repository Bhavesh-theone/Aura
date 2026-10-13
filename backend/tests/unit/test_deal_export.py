"""Deal export: bundle assembly (sections, masking, walls), rupee formatting, renderers."""
import io
import json
import zipfile
from datetime import datetime, timezone
from types import SimpleNamespace as NS

import pytest

from aura.core.domain.deal import Deal, DealSignoff, StageHistoryEntry
from aura.core.domain.diligence import ConditionPrecedent, DiligenceItem, Finding
from aura.core.domain.document import Document
from aura.core.domain.enums import (DealStage, DiligenceSeverity, DiligenceStatus, DocClass,
                                    FindingType, IdentifierScheme, PartyRole, PartyType, Role,
                                    TransactionClass)
from aura.core.domain.errors import Forbidden, NotFound, RuleViolation
from aura.core.domain.export import SECTIONS
from aura.core.domain.party import DealParty, Party, PartyIdentifier
from aura.core.domain.wall import WallMember
from aura.core.money import format_inr, format_inr_compact, group_indian, to_inr_int
from aura.core.services.deal_export import (ExportSources, build_deal_export,
                                            deals_list_table, filter_deals, mask_email,
                                            parse_sections)
from aura.exports import deal_filename, render_deal, render_list

T0 = datetime(2026, 10, 13, 9, 30, tzinfo=timezone.utc)


# --------------------------------------------------------------------------- money

def test_indian_grouping_and_lakh_crore():
    assert group_indian(0) == "0"
    assert group_indian(999) == "999"
    assert group_indian(1000) == "1,000"
    assert group_indian(123456) == "1,23,456"
    assert group_indian(12345678) == "1,23,45,678"
    assert group_indian(-1234567) == "-12,34,567"
    assert format_inr(485000000) == "₹48,50,00,000"
    assert format_inr_compact(485000000) == "₹48.50 Cr"
    assert format_inr_compact(420000) == "₹4.20 L"
    assert format_inr_compact(86400) == "₹86,400"
    assert format_inr_compact(-25000000) == "-₹2.50 Cr"
    assert format_inr_compact(None) == "" and format_inr(None) == ""
    assert format_inr_compact(1234567890000) == "₹1,23,456.79 Cr"
    assert to_inr_int("12.5") == 13 and to_inr_int(True) is None and to_inr_int("x") is None


# --------------------------------------------------------------------------- helpers

def test_parse_sections_orders_validates_and_keeps_summary():
    assert parse_sections(None) == list(SECTIONS)
    assert parse_sections("") == list(SECTIONS)
    assert parse_sections("documents, parties") == ["summary", "parties", "documents"]
    assert parse_sections(["audit"]) == ["summary", "audit"]
    with pytest.raises(RuleViolation, match="nonsense"):
        parse_sections("parties,nonsense")


def test_mask_email():
    assert mask_email("priya.shah@kaveri.in") == "p***@kaveri.in"
    assert mask_email("") == "" and mask_email("nodomain") == "***"


# --------------------------------------------------------------------------- fakes

def _deal(**kw) -> Deal:
    base = dict(id=7, external_id="AUR-0002", name="Project Kaveri", stage=DealStage.DILIGENCE,
                transaction_class=TransactionClass.MNA, owner="bhavesh", outcome=None,
                opened_at=T0, regulatory_flags=["CCI"],
                profile={"structure": "share_purchase", "enterprise_value_inr": 4_850_000_000,
                         "stake_pct": 51, "target_listed": False})
    base.update(kw)
    return Deal(**base)


def _doc(i, upsi=False, sensitivity=2):
    return Document(id=i, deal_id=7, doc_class=DocClass.SPA, filename=f"doc{i}.pdf",
                    sha256="ab" * 32, byte_size=1000, current_version=1, contains_upsi=upsi,
                    sensitivity=sensitivity, uploaded_by="bhavesh", uploaded_at=T0)


def sources(deal: Deal | None = None, docs=None, visible=True) -> ExportSources:
    deal = deal or _deal()
    party = Party(id=3, party_type=PartyType.COMPANY, display_name="Kaveri Foods Pvt Ltd",
                  created_at=T0, identifiers=[PartyIdentifier(IdentifierScheme.PAN,
                                                              "ABCDE1234F")])
    items = [DiligenceItem(1, 7, None, "Tax notices", DiligenceSeverity.STANDARD),
             DiligenceItem(2, 7, None, "Title to land", DiligenceSeverity.BLOCKER),
             DiligenceItem(3, 7, None, "Pending litigation", DiligenceSeverity.MATERIAL,
                           DiligenceStatus.SATISFIED)]
    finding = Finding(9, 2, 7, FindingType.RED_FLAG, "Encumbrance on plot 14", 12_500_000, True,
                      "asha", T0)
    return ExportSources(
        deals=NS(get=lambda i: deal if visible and i == deal.id else None),
        history=NS(list_for_deal=lambda i: [StageHistoryEntry(1, 7, None,
                                                              DealStage.ORIGINATED, T0,
                                                              "bhavesh", None, False)]),
        walls=NS(members=lambda i: [WallMember("bhavesh", T0), WallMember("asha", T0)]),
        parties=NS(list_for_deal=lambda i: [DealParty(7, 3, PartyRole.TARGET,
                                                      "Kaveri Foods Pvt Ltd")],
                   get=lambda i: party),
        contacts=NS(list_for_party=lambda i: [NS(name="Ravi", email="ravi@kaveri.in",
                                                 is_primary=True, verified_at=T0)]),
        mandates=NS(list_for_deal=lambda i: []),
        documents=NS(list_for_deal=lambda i: docs if docs is not None else [_doc(1)]),
        diligence=NS(list_items=lambda i: items),
        findings=NS(list_for_deal=lambda i: [finding]),
        conditions=NS(list_for_deal=lambda i: [ConditionPrecedent(1, 7, "CCI approval",
                                                                  "regulatory")]),
        signoffs=NS(list_for_deal=lambda i: [DealSignoff(1, 7, "priya", "Cleared", T0)]),
        extractions=NS(list_for_deal=lambda i: [NS(id=4, payload={"revenue_inr": 912_000_000,
                                                                  "note": "x"},
                                                   schema_name="financials.v1",
                                                   document_id=1, version=1,
                                                   confidence=0.97,
                                                   human_verified_by=None)]),
        handoffs=NS(list_all=lambda: [NS(
            id=5, deal_id=7, created_at=T0, task_type="draft_email",
            trigger=NS(value="policy_required"), status=NS(value="resolved"),
            context_package={"decision_required": {"question": "Send the NDA?"}},
            directive={"decision": "approve", "reasoning": "Agreed with the client."},
            resolved_by="bhavesh", resolved_at=T0), NS(id=6, deal_id=99)]),
        agent_runs=NS(list_for=lambda deal_id, limit: [
            NS(archetype="extractor", model_id="qwen3.5", model_version="4b", tokens_in=100,
               tokens_out=20, latency_ms=900, started_at=T0)] * 3),
        reader=NS(guard_evaluations=lambda i: [], confirmed_intake=lambda i: None),
        audit_status=lambda: {"ok": True, "head_id": 42, "head_hash": "f" * 64},
        profile_schema=lambda cls, v: {"title": "M&A deal profile", "properties": {
            "structure": {"title": "Structure", "enum": ["share_purchase"],
                          "x-enum-labels": {"share_purchase": "Share purchase"}},
            "enterprise_value_inr": {"title": "Enterprise value"},
            "stake_pct": {"title": "Stake (%)"}, "target_listed": {"title": "Target listed"}}})


def _rows(bundle, section, table):
    return next(t for t in bundle.section(section).tables if t.key == table).rows


# --------------------------------------------------------------------------- assembly

def test_admin_is_forbidden_and_off_wall_is_not_found():
    with pytest.raises(Forbidden):
        build_deal_export(7, "admin", Role.ADMIN, None, sources())
    with pytest.raises(NotFound):
        build_deal_export(7, "asha", Role.ANALYST, None, sources(visible=False))


def test_full_bundle_contents():
    b = build_deal_export(7, "bhavesh", Role.DEALMAKER, None, sources(), now=T0)
    assert b.section_keys == list(SECTIONS)
    assert b.deal["external_id"] == "AUR-0002" and b.deal["team"] == ["bhavesh", "asha"]
    assert b.generated_at == "2026-10-13T09:30:00Z"
    # diligence by severity, blocker first; finding impact in whole rupees
    assert [r["severity"] for r in _rows(b, "diligence", "items")] == \
        ["blocker", "material", "standard"]
    assert _rows(b, "diligence", "findings")[0]["impact_inr"] == 12_500_000
    # profile rendered with the schema's titles and enum labels
    terms = {r["field"]: r for r in _rows(b, "profile", "terms")}
    assert terms["Structure"]["value"] == "Share purchase"
    assert terms["Enterprise value"]["amount_inr"] == 4_850_000_000
    assert terms["Enterprise value"]["value"] == "₹485.00 Cr"
    assert terms["Target listed"]["value"] == "No"
    # financials: extraction figures (numbers only) and profile amounts, with sources
    fin = _rows(b, "financials", "figures")
    assert {(r["metric"], r["amount_inr"]) for r in fin} >= {
        ("Revenue from operations", 912_000_000), ("Enterprise value", 4_850_000_000)}
    assert any(r["source"].startswith("Extraction #4") for r in fin)
    # decisions only for this deal, with the question and the reasoning
    decisions = _rows(b, "decisions", "handoffs")
    assert len(decisions) == 1 and decisions[0]["question"] == "Send the NDA?"
    assert decisions[0]["reasoning"] == "Agreed with the client."
    runs = _rows(b, "agent_runs", "by_model")
    assert runs == [{"archetype": "extractor", "model": "qwen3.5 4b", "runs": 3,
                     "tokens_in": 300, "tokens_out": 60, "first_at": "2026-10-13T09:30:00Z",
                     "last_at": "2026-10-13T09:30:00Z", "avg_latency_ms": 900}]
    assert b.audit["head_id"] == 42 and b.audit["ok"] is True
    # dealmakers see contact emails in full
    assert _rows(b, "parties", "contacts")[0]["email"] == "ravi@kaveri.in"
    assert not b.upsi


def test_analyst_gets_masked_emails_and_section_selection():
    b = build_deal_export(7, "asha", Role.ANALYST, "parties,documents", sources())
    assert b.section_keys == ["summary", "parties", "documents"]
    contacts = next(t for t in b.section("parties").tables if t.key == "contacts")
    assert contacts.rows[0]["email"] == "r***@kaveri.in" and "masked" in contacts.note
    compliance = build_deal_export(7, "priya", Role.COMPLIANCE, "parties", sources())
    assert _rows(compliance, "parties", "contacts")[0]["email"] == "ravi@kaveri.in"


def test_upsi_from_flag_or_sensitivity_four():
    assert build_deal_export(7, "a", Role.ANALYST, "documents",
                             sources(docs=[_doc(1, upsi=True)])).upsi
    assert build_deal_export(7, "a", Role.ANALYST, "documents",
                             sources(docs=[_doc(1, sensitivity=4)])).upsi


def test_json_dict_is_stable_and_machine_typed():
    b = build_deal_export(7, "bhavesh", Role.DEALMAKER, "diligence,profile", sources(), now=T0)
    j = b.to_json_dict()
    assert j["schema_version"] == 1 and set(j["sections"]) == {"summary", "profile",
                                                               "diligence"}
    finding = j["sections"]["diligence"]["tables"]["findings"]["rows"][0]
    assert finding["impact_inr"] == 12_500_000 and finding["raised_at"].endswith("Z")
    assert j["sections"]["profile"]["profile"]["enterprise_value_inr"] == 4_850_000_000
    json.dumps(j)  # serialisable as-is


def test_deals_list_filters_and_table():
    deals = [_deal(), _deal(id=8, external_id="AUR-102", name="Nilgiri Debt",
                            transaction_class=TransactionClass.DEBT, stage=DealStage.CLOSED,
                            profile={"amount_inr": 2_000_000_000})]
    assert [d.id for d in filter_deals(deals, q="nilgiri")] == [8]
    assert [d.id for d in filter_deals(deals, classes=["mna"])] == [7]
    assert [d.id for d in filter_deals(deals, outcomes=["OPEN"])] == [7, 8]
    assert [d.id for d in filter_deals(deals, stages=["CLOSED"])] == [8]
    table = deals_list_table(deals, Role.ANALYST)
    assert [r["headline_inr"] for r in table.rows] == [4_850_000_000, 2_000_000_000]
    with pytest.raises(Forbidden):
        deals_list_table(deals, Role.ADMIN)


# --------------------------------------------------------------------------- renderers

@pytest.fixture(scope="module")
def bundle():
    return build_deal_export(7, "bhavesh", Role.DEALMAKER, None,
                             sources(docs=[_doc(1, upsi=True)]), now=T0)


def test_filename():
    assert deal_filename("AUR-0002", "Project Kaveri", "2026-10-13T09:30:00Z", "pdf") == \
        "AUR-0002_Project-Kaveri_2026-10-13.pdf"


def test_render_pdf_has_text_rupees_and_watermark(bundle):
    import pymupdf
    out = render_deal(bundle, "pdf")
    assert out.media_type == "application/pdf" and out.content[:5] == b"%PDF-"
    doc = pymupdf.open(stream=out.content, filetype="pdf")
    text = "".join(page.get_text() for page in doc)
    assert "AUR-0002" in text and "Project Kaveri" in text
    assert "₹485.00 Cr" in text and "₹1.25 Cr" in text
    assert "UPSI" in text and "Page 1 of" in text


def test_render_docx(bundle):
    from docx import Document as Docx
    out = render_deal(bundle, "docx")
    doc = Docx(io.BytesIO(out.content))
    body = "\n".join(p.text for p in doc.paragraphs)
    cells = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Project Kaveri" in body and "₹1.25 Cr" in cells
    assert "AUR-0002" in doc.sections[0].header.paragraphs[0].text


def test_render_xlsx_typed_cells(bundle):
    from openpyxl import load_workbook
    out = render_deal(bundle, "xlsx")
    wb = load_workbook(io.BytesIO(out.content))
    assert wb.sheetnames[0] == "Summary" and "Diligence and findings" in wb.sheetnames
    ws = wb["Diligence and findings"]
    cells = [c for row in ws.iter_rows() for c in row if c.value == 12_500_000]
    assert cells and "₹" in cells[0].number_format
    fin = wb["Financials"]
    assert fin.freeze_panes == "A3" and fin.auto_filter.ref


def test_render_json_and_zip(bundle):
    j = json.loads(render_deal(bundle, "json").content)
    assert j["deal"]["external_id"] == "AUR-0002" and j["upsi"] is True
    z = zipfile.ZipFile(io.BytesIO(render_deal(bundle, "zip").content))
    names = z.namelist()
    assert "manifest.json" in names and "deal.json" in names
    findings = next(n for n in names if n.endswith("diligence__findings.csv"))
    text = z.read(findings).decode("utf-8-sig")
    assert text.splitlines()[0].startswith("type,description,impact_inr")
    assert "12500000" in text


def test_render_list_csv_and_xlsx():
    table = deals_list_table([_deal()], Role.DEALMAKER)
    csv_bytes = render_list(table, "csv", "2026-10-13T09:30:00Z", "bhavesh").content
    lines = csv_bytes.decode("utf-8-sig").splitlines()
    assert lines[0].startswith("external_id,name") and "4850000000" in lines[1]
    xlsx = render_list(table, "xlsx", "2026-10-13T09:30:00Z", "bhavesh")
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(xlsx.content))["Deals"]
    assert ws["A2"].value == "Reference" and ws["A3"].value == "AUR-0002"
