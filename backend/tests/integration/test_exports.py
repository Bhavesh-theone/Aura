"""Deal exports through the API: every format opens, walls and roles hold, each export is
audited with the sha256 of the exact bytes, and the per-user rate limit answers 429."""
import hashlib
import io
import json
import zipfile

import pytest
from sqlalchemy import text

from aura.config import get_settings
from tests.integration.helpers import link, party, upload

DEAL = {"external_id": "AUR-0002", "name": "Project Kaveri", "transaction_class": "MNA"}


def _deal(api) -> int:
    r = api.call("POST", "/api/deals", json=DEAL)
    assert r.status_code == 201, r.text
    deal_id = r.json()["id"]
    r = api.call("PATCH", f"/api/deals/{deal_id}",
                 json={"profile": {"structure": "share_purchase",
                                   "enterprise_value_inr": 4_850_000_000}})
    assert r.status_code == 200, r.text
    p = party(api, "Kaveri Foods Pvt Ltd")
    link(api, deal_id, p["id"], "target")
    r = api.call("POST", f"/api/parties/{p['id']}/contacts",
                 json={"name": "Ravi", "email": "ravi@kaveri.in"})
    assert r.status_code in (200, 201), r.text
    upload(api, deal_id, "spa.pdf", "spa", content=b"%PDF-1.4 spa")
    api.session.execute(text(
        "INSERT INTO diligence_item (deal_id, prompt, severity, status, evidence_doc_ids) "
        "VALUES (:d, 'Title to land', 'blocker', 'open', '[]')"), {"d": deal_id})
    api.session.commit()
    item = api.session.execute(text("SELECT MAX(id) FROM diligence_item")).scalar()
    r = api.call("POST", f"/api/deals/{deal_id}/findings",
                 json={"diligence_item_id": item, "finding_type": "red_flag",
                       "description": "Encumbrance on plot 14",
                       "quantified_impact_inr": 12_500_000, "affects_valuation": True})
    assert r.status_code == 201, r.text
    return deal_id


def _export(api, deal_id, fmt, user="bhavesh", **params):
    return api.call("GET", f"/api/deals/{deal_id}/export", user=user,
                    params={"format": fmt, **params})


def _audit_rows(api):
    return api.session.execute(text(
        "SELECT deal_id, actor, payload FROM audit_entries WHERE action = 'deal_exported' "
        "ORDER BY id")).all()


def test_pdf_export_opens_and_is_audited_with_its_sha256(api):
    import pymupdf
    deal_id = _deal(api)
    r = _export(api, deal_id, "pdf")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    disposition = r.headers["content-disposition"]
    assert 'filename="AUR-0002_Project-Kaveri_' in disposition and ".pdf" in disposition
    doc = pymupdf.open(stream=r.content, filetype="pdf")
    body = "".join(page.get_text() for page in doc)
    assert "AUR-0002" in body and "Project Kaveri" in body
    assert "₹485.00 Cr" in body and "₹1.25 Cr" in body and "Page 1 of" in body
    sha = hashlib.sha256(r.content).hexdigest()
    assert r.headers["x-content-sha256"] == sha
    rows = _audit_rows(api)
    assert len(rows) == 1
    deal, actor, payload = rows[0]
    assert (deal, actor) == (deal_id, "bhavesh")
    assert payload["sha256"] == sha and payload["bytes"] == len(r.content)
    assert payload["format"] == "pdf" and payload["sections"][0] == "summary"


def test_docx_xlsx_json_zip_open(api):
    from docx import Document
    from openpyxl import load_workbook
    deal_id = _deal(api)

    r = _export(api, deal_id, "docx")
    assert r.status_code == 200 and r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml")
    doc = Document(io.BytesIO(r.content))
    assert "Project Kaveri" in "\n".join(p.text for p in doc.paragraphs)

    r = _export(api, deal_id, "xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "Diligence and findings" in wb.sheetnames
    amounts = [c for row in wb["Diligence and findings"].iter_rows() for c in row
               if c.value == 12_500_000]
    assert amounts and "₹" in amounts[0].number_format

    r = _export(api, deal_id, "json", sections="parties,diligence")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/json")
    body = r.json()
    assert body["schema_version"] == 1 and body["deal"]["external_id"] == "AUR-0002"
    assert list(body["sections"]) == ["summary", "parties", "diligence"]
    finding = body["sections"]["diligence"]["tables"]["findings"]["rows"][0]
    assert finding["impact_inr"] == 12_500_000
    # dealmakers see contact emails in full
    contacts = body["sections"]["parties"]["tables"]["contacts"]["rows"]
    assert contacts[0]["email"] == "ravi@kaveri.in"
    assert body["audit"]["ok"] is True and len(body["audit"]["head_hash"]) == 64

    r = _export(api, deal_id, "zip")
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = z.namelist()
    assert "manifest.json" in names and any(n.endswith("documents__documents.csv")
                                            for n in names)
    manifest = json.loads(z.read("manifest.json"))
    assert manifest["deal"]["external_id"] == "AUR-0002"
    assert len(_audit_rows(api)) == 4


def test_permissions_wall_and_roles(api):
    deal_id = _deal(api)
    # analyst not on the wall: the deal does not exist for her
    assert _export(api, deal_id, "json", user="asha").status_code == 404
    # admin has no deal powers at all
    assert _export(api, deal_id, "json", user="admin").status_code == 403
    assert api.call("GET", "/api/deals/export", user="admin").status_code == 403
    api.add_member(deal_id, "asha")
    r = _export(api, deal_id, "json", user="asha", sections="parties")
    assert r.status_code == 200
    contacts = r.json()["sections"]["parties"]["tables"]["contacts"]["rows"]
    assert contacts[0]["email"] == "r***@kaveri.in"
    # compliance (on every new deal's wall) may export too
    assert _export(api, deal_id, "pdf", user="priya").status_code == 200
    # bad section names and formats are refused
    assert _export(api, deal_id, "json", sections="nonsense").status_code == 422
    assert _export(api, deal_id, "exe").status_code == 422
    assert _export(api, 999999, "json").status_code == 404
    # nothing was audited for refused exports
    assert [r[1] for r in _audit_rows(api)] == ["asha", "priya"]


def test_deals_list_export_csv_xlsx_with_filters_and_walls(api):
    from openpyxl import load_workbook
    deal_id = _deal(api)
    r = api.call("POST", "/api/deals", json={"external_id": "AUR-102", "name": "Nilgiri Debt",
                                             "transaction_class": "DEBT"})
    assert r.status_code == 201
    r = api.call("GET", "/api/deals/export", params={"format": "csv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert 'filename="Aura_deals_' in r.headers["content-disposition"]
    lines = r.content.decode("utf-8-sig").splitlines()
    assert lines[0].startswith("external_id,name,transaction_class") and len(lines) == 3
    r = api.call("GET", "/api/deals/export",
                 params={"format": "xlsx", "transaction_class": ["MNA"], "q": "kaveri"})
    ws = load_workbook(io.BytesIO(r.content))["Deals"]
    assert ws["A3"].value == "AUR-0002" and ws["A4"].value is None
    assert ws["J3"].value == 4_850_000_000
    # the analyst sees only deals on her wall
    api.add_member(deal_id, "asha")
    r = api.call("GET", "/api/deals/export", user="asha", params={"format": "csv"})
    lines = r.content.decode("utf-8-sig").splitlines()
    assert len(lines) == 2 and lines[1].startswith("AUR-0002")
    payload = _audit_rows(api)[-1][2]
    assert payload["scope"] == "deals_list" and payload["rows"] == 1
    assert payload["sha256"] == hashlib.sha256(r.content).hexdigest()


def test_export_rate_limit(api, monkeypatch):
    from aura.api import rate_limit
    deal_id = _deal(api)
    monkeypatch.setattr(get_settings(), "exports_per_minute", 2)
    rate_limit.reset()
    assert _export(api, deal_id, "json").status_code == 200
    assert api.call("GET", "/api/deals/export").status_code == 200
    r = _export(api, deal_id, "json")
    assert r.status_code == 429 and int(r.headers["retry-after"]) >= 1
    # another user has their own bucket
    assert _export(api, deal_id, "json", user="priya").status_code == 200


@pytest.mark.parametrize("fmt", ["pdf", "docx", "xlsx", "json", "zip"])
def test_upsi_deal_marks_every_human_export(api, fmt):
    deal_id = _deal(api)
    r = api.call("POST", f"/api/deals/{deal_id}/documents",
                 files={"file": ("board.pdf", b"%PDF-1.4 upsi")},
                 data={"doc_class": "board_resolution", "contains_upsi": "true"})
    assert r.status_code == 201, r.text
    r = _export(api, deal_id, fmt)
    assert r.status_code == 200
    if fmt == "pdf":
        import pymupdf
        text_ = "".join(p.get_text() for p in pymupdf.open(stream=r.content, filetype="pdf"))
        assert "UPSI" in text_
    elif fmt == "json":
        assert r.json()["upsi"] is True
    assert _audit_rows(api)[-1][2]["upsi"] is True
